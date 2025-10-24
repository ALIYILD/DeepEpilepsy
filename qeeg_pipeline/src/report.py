"""Report generation helpers using python-docx."""
from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable, Mapping

from docx import Document
from docx.shared import Inches
import pandas as pd

from .config import PatientMetadata


def _add_table(doc: Document, dataframe: pd.DataFrame, title: str) -> None:
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=len(dataframe.columns))
    hdr = table.rows[0].cells
    for idx, column in enumerate(dataframe.columns):
        hdr[idx].text = str(column)
    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = f"{value}"


def build_report(
    metadata: PatientMetadata,
    summary_rows: Iterable[Mapping[str, str | float]],
    figure_paths: Iterable[Path],
    asymmetry_table: pd.DataFrame,
    pdr_detail: pd.DataFrame,
    output_path: Path,
) -> Path:
    doc = Document()
    doc.add_heading(f"QEEG Report – {metadata.client_id}", 0)

    demographics = []
    if metadata.name:
        demographics.append(f"Name: {metadata.name}")
    if metadata.date_of_birth:
        demographics.append(f"DOB: {metadata.date_of_birth}")
    if metadata.recording_date:
        demographics.append(f"Recorded: {metadata.recording_date}")
    if metadata.handedness:
        demographics.append(f"Handedness: {metadata.handedness}")
    if metadata.medications:
        demographics.append("Medications: " + ", ".join(metadata.medications))
    if metadata.notes:
        demographics.append(f"Notes: {metadata.notes}")

    if demographics:
        doc.add_paragraph(" | ".join(demographics))

    doc.add_heading("Summary Indices", level=1)
    summary_df = pd.DataFrame(summary_rows)
    _add_table(doc, summary_df, title="Global Metrics")

    if not asymmetry_table.empty:
        _add_table(doc, asymmetry_table, title="Alpha Asymmetry (log power)")
    if not pdr_detail.empty:
        _add_table(doc, pdr_detail, title="Posterior Dominant Rhythm (per channel)")

    doc.add_heading("Figures", level=1)
    for path in figure_paths:
        if path.exists():
            doc.add_paragraph(path.stem.replace("_", " "))
            doc.add_picture(str(path), width=Inches(4))

    doc.add_heading("Key Findings", level=1)
    doc.add_paragraph(
        "This automatically generated report summarizes quantitative EEG metrics. "
        "Please integrate these data with clinical history and other assessments "
        "before making diagnostic or therapeutic decisions."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
